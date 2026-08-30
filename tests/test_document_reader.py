from docx import Document

from src.document_reader import read_docx


def test_read_paragraph_and_table(tmp_path):
    path = tmp_path / "x.docx"
    doc = Document()
    doc.add_paragraph("  测试段落  ")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "时间"
    table.cell(0, 1).text = "周三"
    doc.save(path)
    result = read_docx(path)
    assert result[0]["text"] == "测试段落"
    assert "时间；周三" == result[1]["text"]
