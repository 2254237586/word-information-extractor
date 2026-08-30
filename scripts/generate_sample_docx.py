from pathlib import Path

from docx import Document


def main() -> None:
    output = Path(__file__).resolve().parents[1] / "data" / "sample.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for text in (
        "关于本周课程与实验安排的通知",
        "各位同学，本周教学安排如下。",
        "机械原理实验将于周三下午2点在综合楼305进行，请大家提前十分钟到场。",
        "实验报告需于本周日22:00前上传至学习通。",
        "今天的天气不错，大家注意休息。",
        "四六级考试报名将于8月28日下午5点截止，请还没有报名的同学及时完成报名。",
        "下周一上午10点将在研教楼203进行课程答疑。",
        "最近课程进度比较快，大家课后可以适当复习。",
        "本周五下午将组织实验室安全培训，地点为机械馆201。",
        "请参加培训的同学提前签到。",
    ):
        doc.add_paragraph(text)
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ("事项", "时间", "地点")):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ("课程答疑", "下周一上午10点", "研教楼203")):
        cell.text = value
    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()

