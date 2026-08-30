import re
from pathlib import Path
from typing import Dict, List, Union

from docx import Document


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def read_docx(file_path: Union[str, Path]) -> List[Dict[str, object]]:
    """读取段落和表格文字，并保留稳定、递增的来源编号。"""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Word 文件不存在：{path}")
    if path.suffix.lower() != ".docx":
        raise ValueError("仅支持 .docx 格式的 Word 文件")
    try:
        document = Document(path)
    except Exception as exc:
        raise ValueError(f"无法读取 Word 文件，文件可能已损坏：{path}") from exc

    items: List[Dict[str, object]] = []
    for paragraph in document.paragraphs:
        text = _normalize(paragraph.text)
        if text:
            items.append({"paragraph_id": len(items), "text": text, "source_type": "paragraph"})
    for table_id, table in enumerate(document.tables):
        for row_id, row in enumerate(table.rows):
            cells = [_normalize(cell.text) for cell in row.cells]
            text = "；".join(dict.fromkeys(cell for cell in cells if cell))
            if text:
                items.append({"paragraph_id": len(items), "text": text,
                              "source_type": "table", "table_id": table_id, "row_id": row_id})
    if not items:
        raise ValueError("Word 文档中没有可提取的文字")
    return items

